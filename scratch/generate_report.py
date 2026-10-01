import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_report():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    def add_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return h

    def add_paragraph(text, bold=False, italic=False, align='justify'):
        p = doc.add_paragraph(text)
        if bold: p.runs[0].bold = True
        if italic: p.runs[0].italic = True
        if align == 'justify':
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    # --- COVER PAGE ---
    doc.add_paragraph("\n" * 2)
    p = doc.add_paragraph("A PROJECT REPORT ON")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(16)
    
    p = doc.add_paragraph("COLLEGE ERP SYSTEM WITH SECURE PROCTORED EXAMINATION MODULE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(22)
    
    doc.add_paragraph("\n" * 2)
    
    # Image Placeholder for Cover
    try:
        doc.add_picture(r'C:\Users\manid\.gemini\antigravity\brain\fe17b431-23f8-4ef4-bd2c-51598c930831\erp_dashboard_preview_1777826319420.png', width=Inches(5))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except:
        pass

    doc.add_paragraph("\n" * 2)
    
    p = doc.add_paragraph("SUBMITTED IN PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE DEGREE OF")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph("BACHELOR OF ENGINEERING")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    
    doc.add_paragraph("\n" * 1)
    
    p = doc.add_paragraph("SUBMITTED BY:")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("[STUDENT NAME 1] ([ROLL NO])\n[STUDENT NAME 2] ([ROLL NO])")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    doc.add_paragraph("\n" * 1)
    
    p = doc.add_paragraph("UNDER THE GUIDANCE OF:")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("[GUIDE NAME]\nAssistant Professor")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    doc.add_paragraph("\n" * 2)
    
    p = doc.add_paragraph("DEPARTMENT OF COMPUTER ENGINEERING")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p = doc.add_paragraph("[COLLEGE NAME]\n[UNIVERSITY NAME]\n2025-2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # --- CERTIFICATE ---
    add_heading("CERTIFICATE", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("\n")
    cert_text = (
        "This is to certify that the project entitled \"COLLEGE ERP SYSTEM WITH SECURE PROCTORED EXAMINATION MODULE\" "
        "is a bonafide work carried out by [STUDENT NAME 1] and [STUDENT NAME 2] in partial fulfillment of the requirements "
        "for the degree of Bachelor of Engineering in Computer Engineering of [UNIVERSITY NAME] during the academic year 2025-2026."
    )
    add_paragraph(cert_text)
    doc.add_paragraph("\n" * 4)
    
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "Guide\n[GUIDE NAME]"
    table.cell(0, 1).text = "Head of Department\n[HOD NAME]"
    table.cell(0, 2).text = "Principal\n[PRINCIPAL NAME]"
    
    doc.add_page_break()

    # --- ABSTRACT ---
    add_heading("ABSTRACT", level=1)
    abstract_text = (
        "The College ERP (Enterprise Resource Planning) System is an integrated software solution designed to manage and "
        "automate various academic and administrative operations within an educational institution. This project focuses "
        "on building a robust, scalable, and secure platform that facilitates seamless interaction between students, "
        "faculty, and administration. A key highlight of this system is the 'Secure Proctored Examination Module', which "
        "utilizes WebSockets and real-time event tracking to prevent academic dishonesty during online assessments. "
        "The backend is engineered using Java Spring Boot with a Microservices-ready modular architecture, ensuring high "
        "availability and performance. The frontend is developed using React.js for a dynamic and responsive user experience. "
        "Security is handled via JWT-based stateless authentication, and Firebase is utilized for scalable cloud storage. "
        "The system also integrates automated email notifications, attendance tracking, and comprehensive dashboard analytics."
    )
    add_paragraph(abstract_text)
    doc.add_page_break()

    # --- TABLE OF CONTENTS (Placeholder) ---
    add_heading("TABLE OF CONTENTS", level=1)
    toc_items = [
        "1. Introduction",
        "2. System Architecture & Tech Stack",
        "3. Backend Engineering Deep Dive",
        "4. Database Design",
        "5. API Documentation",
        "6. Implementation Details",
        "7. Deployment & CI/CD",
        "8. Conclusion & Future Enhancements",
        "9. References"
    ]
    for item in toc_items:
        doc.add_paragraph(item)
    doc.add_page_break()

    # --- CHAPTER 1: INTRODUCTION ---
    add_heading("1. INTRODUCTION", level=1)
    add_paragraph("1.1 Project Overview", bold=True)
    add_paragraph(
        "In the modern era of digital education, managing institutional data manually is both inefficient and error-prone. "
        "The College ERP System aims to bridge this gap by providing a centralized platform for all academic activities. "
        "From student enrollment to examination management, the system provides a comprehensive suite of tools."
    )
    add_paragraph("1.2 Problem Statement", bold=True)
    add_paragraph(
        "Existing systems often lack real-time monitoring during online exams, leading to integrity issues. "
        "Furthermore, fragmented data across different departments makes administration challenging. "
        "This project addresses these issues by providing a secure proctoring environment and a unified data management system."
    )

    # --- CHAPTER 2: SYSTEM ARCHITECTURE ---
    add_heading("2. SYSTEM ARCHITECTURE & TECH STACK", level=1)
    add_paragraph("2.1 Backend Architecture", bold=True)
    add_paragraph(
        "The backend is built on a Layered Architecture (Controller-Service-Repository) using Java Spring Boot. "
        "This separation of concerns ensures maintainability and scalability."
    )
    try:
        doc.add_picture(r'C:\Users\manid\.gemini\antigravity\brain\fe17b431-23f8-4ef4-bd2c-51598c930831\backend_architecture_diagram_1777825902652.png', width=Inches(5))
        p = doc.add_paragraph("Figure 2.1: System Architecture Overview")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except:
        pass

    add_paragraph("2.2 Technology Stack", bold=True)
    tech_table = doc.add_table(rows=1, cols=2)
    tech_table.style = 'Table Grid'
    hdr_cells = tech_table.rows[0].cells
    hdr_cells[0].text = 'Layer'
    hdr_cells[1].text = 'Technology'
    
    data = [
        ('Frontend', 'React.js, Redux, TailwindCSS'),
        ('Backend', 'Java Spring Boot, Spring Security'),
        ('Database', 'MySQL / PostgreSQL'),
        ('Cloud Storage', 'Google Firebase'),
        ('Authentication', 'JWT (JSON Web Tokens)'),
        ('Communication', 'REST API, WebSockets (STOMP)'),
        ('Build Tools', 'Maven, Vite')
    ]
    for layer, tech in data:
        row_cells = tech_table.add_row().cells
        row_cells[0].text = layer
        row_cells[1].text = tech

    # --- CHAPTER 3: BACKEND ENGINEERING ---
    add_heading("3. BACKEND ENGINEERING DEEP DIVE", level=1)
    add_paragraph("3.1 Request-Response Lifecycle", bold=True)
    add_paragraph(
        "Every API request follows a strict lifecycle: Authentication Filter -> Controller Mapping -> Service Logic -> Repository Data Access -> DTO Mapping -> Response."
    )
    add_paragraph("3.2 Authentication Flow (JWT)", bold=True)
    add_paragraph(
        "We implement stateless authentication using JWT. Upon login, the server generates a token which is stored in the client-side (HttpOnly cookies or LocalStorage). "
        "Subsequent requests include this token in the 'Authorization' header."
    )
    add_paragraph("3.3 Real-time Communication via WebSockets", bold=True)
    add_paragraph(
        "For the proctoring module, WebSockets (using STOMP protocol) are utilized to transmit violation alerts from the student's browser to the teacher's dashboard in under 100ms."
    )

    # --- CHAPTER 4: DATABASE DESIGN ---
    add_heading("4. DATABASE DESIGN", level=1)
    try:
        doc.add_picture(r'C:\Users\manid\.gemini\antigravity\brain\fe17b431-23f8-4ef4-bd2c-51598c930831\database_er_diagram_1777826507879.png', width=Inches(5))
        p = doc.add_paragraph("Figure 4.1: Database ER Diagram")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except:
        pass
    
    add_paragraph("4.1 Schema Overview", bold=True)
    add_paragraph(
        "The database follows 3NF (Third Normal Form) to minimize redundancy. Primary tables include 'Users', 'Students', 'Professors', 'Courses', 'Exams', and 'Violations'."
    )

    # --- CHAPTER 5: API DOCUMENTATION ---
    add_heading("5. API DOCUMENTATION", level=1)
    api_table = doc.add_table(rows=1, cols=4)
    api_table.style = 'Table Grid'
    hdr = api_table.rows[0].cells
    hdr[0].text = 'Endpoint'
    hdr[1].text = 'Method'
    hdr[2].text = 'Description'
    hdr[3].text = 'Auth'
    
    apis = [
        ('/api/auth/login', 'POST', 'User Authentication', 'None'),
        ('/api/exams/start', 'GET', 'Fetch Exam Questions', 'Student'),
        ('/api/proctor/report', 'WS', 'Real-time Violation Report', 'Student'),
        ('/api/admin/users', 'GET', 'Manage User Accounts', 'Admin'),
        ('/api/attendance/save', 'POST', 'Save Attendance', 'Professor')
    ]
    for ep, meth, desc, auth in apis:
        row = api_table.add_row().cells
        row[0].text = ep
        row[1].text = meth
        row[2].text = desc
        row[3].text = auth

    # --- CHAPTER 6: IMPLEMENTATION DETAILS ---
    add_heading("6. IMPLEMENTATION DETAILS", level=1)
    add_paragraph("6.1 Backend Folder Structure", bold=True)
    structure = (
        "src/main/java/com/example/stud_erp/\n"
        "├── controller/       # API Endpoints\n"
        "├── service/          # Business Logic\n"
        "├── repository/       # Database Interfacing\n"
        "├── entity/           # JPA Models\n"
        "├── security/         # JWT & Security Configuration\n"
        "├── payload/          # DTOs & Request/Response objects\n"
        "└── configuration/    # App Configs (WebSocket, CORS)"
    )
    doc.add_paragraph(structure, style='No Spacing').runs[0].font.name = 'Courier New'

    add_paragraph("6.2 Error Handling Strategy", bold=True)
    add_paragraph(
        "Global exception handling is implemented using '@ControllerAdvice', ensuring that the API returns consistent JSON error responses across all failure scenarios."
    )

    # --- CHAPTER 7: DEPLOYMENT ---
    add_heading("7. DEPLOYMENT & CI/CD", level=1)
    add_paragraph("7.1 Deployment Architecture", bold=True)
    add_paragraph(
        "The application is containerized using Docker and deployed on AWS EC2 instances. Static assets (React Build) are served via Nginx."
    )
    add_paragraph("7.2 CI/CD Pipeline", bold=True)
    add_paragraph(
        "GitHub Actions is used for CI/CD. Every push to the 'main' branch triggers automated tests and subsequent deployment to the staging environment."
    )

    # --- CHAPTER 8: CONCLUSION ---
    add_heading("8. CONCLUSION & FUTURE ENHANCEMENTS", level=1)
    add_paragraph(
        "The College ERP System successfully automates major administrative tasks and provides a secure environment for online examinations. "
        "Future enhancements include AI-based facial recognition for proctoring and mobile application integration for better accessibility."
    )

    # --- REFERENCES ---
    add_heading("REFERENCES", level=1)
    refs = [
        "1. Spring Boot Documentation - https://spring.io/projects/spring-boot",
        "2. React Documentation - https://reactjs.org/",
        "3. JWT Guide - https://jwt.io/introduction/",
        "4. WebSocket STOMP Protocol - https://stomp.github.io/"
    ]
    for ref in refs:
        doc.add_paragraph(ref)

    # Save document
    output_path = r'd:\Reactjs\Erp-website\College-ERP-Using-Reactjs-And-Java-Spring-Boot\College_ERP_Project_Report.docx'
    doc.save(output_path)
    print(f"Report saved to {output_path}")

if __name__ == "__main__":
    create_report()
